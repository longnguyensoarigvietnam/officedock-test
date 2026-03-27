from datetime import datetime, timedelta
from io import BytesIO, StringIO
from pathlib import Path
import csv
import re
from collections import defaultdict

from django.db.models import Case, When
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Side

from users.models import User
from statistics.constants import ExportType, PeriodClassification
from tags.models import Tag
from organizations.models import Organization
from skills.models import StatisticCategory
from stat_data.constants import ALL_TEAM
from stat_data.utils import (
    normalize_percentages,
)
from common.utils import format_duration, time_str_to_timedelta


class ExportTaskService:
    """
    Service for exporting task statistics to CSV or Excel files.

    This class loads a predefined template, fills or processes it (if needed),
    and then returns the resulting file as an in-memory object for download.
    """

    XLSX_TEMPLATE_PATH = Path("templates/export/tasks.xlsx")
    CSV_TEMPLATE_PATH = Path("templates/export/tasks.csv")

    def __init__(
        self, request, queryset, export_type, sum_total_duration, users=None
    ):
        """
        Initialize the export service.

        Args:
            request: Django request object.
            queryset: The queryset of task records to be exported.
            export_type: File format for export (CSV or XLSX).
        """
        self.request = request
        self.queryset = queryset
        self.export_type = export_type
        self.sum_total_duration = sum_total_duration
        self.users = users

        # Parse query params once
        q = request.query_params
        self.from_date = q.get("from_date")
        self.end_date = q.get("end_date")
        self.large_category_id = q.get("large_category_id")
        self.medium_category_id = q.get("medium_category_id")
        self.small_category_id = q.get("small_category_id")
        self.tag_ids = [
            t for t in q.get("tag_ids", "").split(",") if t and str(t).isdigit()
        ]
        self.user_id = q.get("user_id")
        self.organization_ids = q.get("organization_ids")
        self.period = q.get("period_classification")

        # Preload all needed info
        self.user = self._get_user()
        self.organization_name = self._get_organization_name()
        self.full_name = getattr(self.user, "full_name", "")
        self.period_classification = (
            PeriodClassification.COMPARISON.value
            if self.period == PeriodClassification.COMPARISON.name
            else PeriodClassification.BASE.value
        )
        self.tag_names_filter = self._get_tag_filter_names()
        self.category_names = self._get_category_names()

    # ------------------------------------------------------------------ #
    # Internal helpers
    # ------------------------------------------------------------------ #
    def _get_organization_name(self):
        if self.organization_ids == ALL_TEAM:
            return ALL_TEAM

        if self.organization_ids and str(self.organization_ids).isdigit():
            org = Organization.objects.filter(id=self.organization_ids).first()
            return org.name if org else ""

        return ""

    def _get_user(self):
        if self.user_id and str(self.user_id).isdigit():
            return (
                User.objects.filter(id=self.user_id).first()
                or self.request.user
            )
        return self.request.user

    def _get_tag_filter_names(self):
        if not self.tag_ids:
            return ""

        preserved = Case(
            *[When(id=pk, then=pos) for pos, pk in enumerate(self.tag_ids)]
        )

        tag_names = (
            Tag.objects.filter(id__in=self.tag_ids)
            .order_by(preserved)
            .values_list("name", flat=True)
        )

        return " ".join(f"#{name}" for name in tag_names)

    def _get_category_names(self):
        category_ids = [
            cid
            for cid in [
                self.large_category_id,
                self.medium_category_id,
                self.small_category_id,
            ]
            if cid and str(cid).isdigit()
        ]
        if not category_ids:
            return {"large": "-", "medium": "-", "small": "-"}
        cats = StatisticCategory.objects.filter(
            id__in=category_ids
        ).values_list("id", "name")
        cmap = {str(cid): cname for cid, cname in cats}
        return {
            "large": cmap.get(str(self.large_category_id), "-"),
            "medium": cmap.get(str(self.medium_category_id), "-"),
            "small": cmap.get(str(self.small_category_id), "-"),
        }

    def format_date(self, date_str):
        try:
            return datetime.strptime(date_str, "%Y-%m-%d").strftime("%Y/%m/%d")
        except Exception:
            return date_str or ""

    # ------------------------------------------------------------------ #
    # Public
    # ------------------------------------------------------------------ #
    def get_filename(self):
        """
        Generate a timestamped filename for the exported file.
        Example output:
            - 20250810タスク一覧集計.csv
            - 20250810タスク一覧集計.xlsx
        """
        timestamp = datetime.now().strftime("%Y%m%d")
        return f"{timestamp}タスク一覧集計.{self.export_type}"

    def export_task_statistic(self):
        """
        Export task statistics based on the chosen file format.
        Returns:
            BytesIO: File-like object containing the exported data.
        """
        if self.export_type == ExportType.CSV.value:
            return self._export_csv()
        elif self.export_type == ExportType.XLSX.value:
            return self._export_excel()
        raise ValueError(f"Unsupported export type: {self.export_type}")

    # ------------------------------------------------------------------ #
    # CSV Export
    # ------------------------------------------------------------------ #
    def _export_csv(self):
        """
        Export data using a CSV template.
        The template provides the header row and structure.
        """

        # Load header template
        with open(self.CSV_TEMPLATE_PATH, "r", encoding="utf-8") as f:
            reader = csv.reader(f)
            rows = list(reader)

        # Build data rows
        for idx, task in enumerate(self.queryset, 1):
            rows.append(self._build_row(idx, task))

        # Write CSV to memory
        output = StringIO()
        writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL)
        writer.writerows(rows)

        csv_bytes = ("\ufeff" + output.getvalue()).encode("utf-8")

        return BytesIO(csv_bytes)

    # ------------------------------------------------------------------ #
    # Excel Export
    # ------------------------------------------------------------------ #
    def _export_excel(self):
        """
        Export data using an Excel (.xlsx) template.
        """
        from users.serializers import BaseUserProfileSerializer

        target_users = self.users
        if not target_users:
            target_users = [self.user] if self.user else []

        user_groups = defaultdict(list)
        for item in self.queryset:
            user_groups[item["user"]["id"]].append(item)

        num_users = len(target_users)

        wb = load_workbook(self.XLSX_TEMPLATE_PATH)
        ws = wb.active

        if num_users > 1:
            # Multiple users, create sheet for each
            used_sheet_names = set()
            for user in target_users:
                items = user_groups.get(user.id, [])
                user_info = BaseUserProfileSerializer(user).data
                total_duration_user = sum(
                    (
                        time_str_to_timedelta(item["total_duration"])
                        or timedelta(0)
                        for item in items
                    ),
                    timedelta(0),
                )
                # Calculate raw percent for each item based on user's total
                for item in items:
                    duration_sec = time_str_to_timedelta(
                        item["total_duration"]
                    ).total_seconds()
                    total_sec = total_duration_user.total_seconds()
                    item["percent"] = (
                        (duration_sec / total_sec) * 100 if total_sec > 0 else 0
                    )

                # Normalize percentages to sum to exactly 100%
                normalize_percentages(items)

                current_ws = wb.copy_worksheet(ws)
                current_ws.title = self._sanitize_sheet_name(
                    user_info.get("full_name", f"User {user.id}"),
                    used_sheet_names,
                )
                # Fill header info
                current_ws[
                    "B4"
                ] = f"{self.format_date(self.from_date)} - {self.format_date(self.end_date)}"
                current_ws["B5"] = self.period_classification
                current_ws["B6"] = user_info.get("full_name", "")
                current_ws[
                    "B7"
                ] = f"{self.organization_name} > {self.category_names['large']} > {self.category_names['medium']} > {self.category_names['small']}"
                current_ws["B8"] = self.tag_names_filter
                current_ws["B9"] = format_duration(total_duration_user)

                # Fill data rows
                thin_border = Border(
                    left=Side(style="thin", color="000000"),
                    right=Side(style="thin", color="000000"),
                    top=Side(style="thin", color="000000"),
                    bottom=Side(style="thin", color="000000"),
                )
                left_align = Alignment(horizontal="left", vertical="center")
                start_row = 13
                if items:
                    for idx, task in enumerate(items, 1):
                        data = self._build_row(
                            idx, task, user_name=user_info.get("full_name", "")
                        )
                        row_values = [
                            data[6],
                            data[7],
                            data[8],
                            data[9],
                            data[10],
                            data[11],
                            data[12],
                            data[13],
                            data[14],
                        ]
                        for col_idx, value in enumerate(row_values, start=1):
                            cell = current_ws.cell(
                                row=start_row, column=col_idx, value=value
                            )
                            cell.border = thin_border

                            # No.
                            if col_idx == 1:
                                cell.alignment = left_align

                            # 割合
                            if col_idx == 4:
                                cell.number_format = "0.00%"
                                cell.alignment = left_align

                        start_row += 1
            # Remove the original template sheet as it's no longer needed
            wb.remove(ws)
            if wb.sheetnames:
                wb.active = 0
        else:
            # Single user, original logic
            user_info = self.queryset[0]["user"] if self.queryset else {}
            user_name = user_info.get("full_name", self.full_name)

            # Fill header info
            ws[
                "B4"
            ] = f"{self.format_date(self.from_date)} - {self.format_date(self.end_date)}"
            ws["B5"] = self.period_classification
            ws["B6"] = user_name
            ws[
                "B7"
            ] = f"{self.organization_name} > {self.category_names['large']} > {self.category_names['medium']} > {self.category_names['small']}"
            ws["B8"] = self.tag_names_filter
            ws["B9"] = self.sum_total_duration

            # Fill data rows
            thin_border = Border(
                left=Side(style="thin", color="000000"),
                right=Side(style="thin", color="000000"),
                top=Side(style="thin", color="000000"),
                bottom=Side(style="thin", color="000000"),
            )
            left_align = Alignment(horizontal="left", vertical="center")
            start_row = 13
            for idx, task in enumerate(self.queryset, 1):
                data = self._build_row(idx, task, user_name=user_name)
                row_values = [
                    data[6],
                    data[7],
                    data[8],
                    data[9],
                    data[10],
                    data[11],
                    data[12],
                    data[13],
                    data[14],
                ]
                for col_idx, value in enumerate(row_values, start=1):
                    cell = ws.cell(row=start_row, column=col_idx, value=value)
                    cell.border = thin_border

                    # No.
                    if col_idx == 1:
                        cell.alignment = left_align

                    # 割合
                    if col_idx == 4:
                        cell.number_format = "0.00%"
                        cell.alignment = left_align

                start_row += 1

        # Save to BytesIO
        excel_file = BytesIO()
        wb.save(excel_file)
        excel_file.seek(0)
        return excel_file

    def _sanitize_sheet_name(self, name, used_names):
        """
        Sanitize sheet name to follow Excel rules:
        - Max 31 characters
        - No forbidden characters: \ / ? * [ ] :
        - Must be unique
        """
        if not name:
            name = "Sheet"
        # Forbidden chars: \ / ? * [ ] :
        name = re.sub(r"[\\/*?:\[\]]", "", name)
        # Limit to 31 chars
        name = name[:31].strip()
        if not name:
            name = "Sheet"

        # Ensure uniqueness
        original_name = name
        counter = 1
        while name.lower() in [n.lower() for n in used_names]:
            suffix = f"({counter})"
            name = original_name[: 31 - len(suffix)] + suffix
            counter += 1

        used_names.add(name)
        return name

    # ------------------------------------------------------------------ #
    # Shared Row Builder
    # ------------------------------------------------------------------ #
    def _build_row(self, idx, task, user_name=None):
        if user_name is None:
            user_name = self.full_name
        title = task.get("title", "")
        total_duration = task.get("total_duration", "00:00")
        percent = int(task.get("percent", 0)) / 100
        org = task.get("organization", {}) or {}
        organization_name = org.get("name", "") if isinstance(org, dict) else ""

        large_name = medium_name = small_name = ""
        for c in task.get("categories", []) or []:
            if not isinstance(c, dict):
                continue
            ctype = c.get("type")
            if ctype == "LARGE":
                large_name = c.get("name", "")
            elif ctype == "MEDIUM":
                medium_name = c.get("name", "")
            elif ctype == "SMALL":
                small_name = c.get("name", "")

        tag_names = " ".join(
            f"#{t.get('name', '')}"
            for t in task.get("tags", [])
            if isinstance(t, dict)
        ).strip()

        return [
            self.from_date,
            self.end_date,
            self.period_classification,
            user_name,
            f"{self.organization_name} > {self.category_names['large']} > {self.category_names['medium']} > {self.category_names['small']}",
            self.tag_names_filter,
            idx,
            title,
            total_duration,
            percent,
            organization_name,
            large_name,
            medium_name,
            small_name,
            tag_names,
        ]
